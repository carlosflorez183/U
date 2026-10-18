"""Exportación de telemetría ambiental DHT22 a JSON, CSV, Excel y PDF."""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.application.exports import ExportedFile, ExportFormat
from app.domain.telemetry import TelemetrySample

_COLUMNS = (
    "id",
    "group_id",
    "device_id",
    "sequence",
    "measured_at",
    "received_at",
    "sampling_period_seconds",
    "temperature_celsius",
    "humidity_percent",
    "temp_status",
    "humidity_status",
    "quality_flags",
)


class FileTelemetryExportRenderer:
    """Genera archivos con columnas estables para análisis ambiental."""

    def render(
        self, samples: Sequence[TelemetrySample], export_format: ExportFormat
    ) -> ExportedFile:
        rows = [_to_row(sample) for sample in samples]
        renderers = {
            ExportFormat.JSON: self._json,
            ExportFormat.CSV: self._csv,
            ExportFormat.XLSX: self._xlsx,
            ExportFormat.PDF: self._pdf,
        }
        media_types = {
            ExportFormat.JSON: "application/json; charset=utf-8",
            ExportFormat.CSV: "text/csv; charset=utf-8",
            ExportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ExportFormat.PDF: "application/pdf",
        }
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        return ExportedFile(
            content=renderers[export_format](rows),
            filename=f"telemetria_iot_{timestamp}.{export_format.value}",
            media_type=media_types[export_format],
        )

    @staticmethod
    def _json(rows: list[dict]) -> bytes:
        return json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")

    @staticmethod
    def _csv(rows: list[dict]) -> bytes:
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
        return buffer.getvalue().encode("utf-8-sig")

    @staticmethod
    def _xlsx(rows: list[dict]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Telemetría IoT"
        sheet.append(list(_COLUMNS))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E7D32")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:L{max(1, len(rows) + 1)}"
        for row in rows:
            sheet.append([_safe_cell(row[column]) for column in _COLUMNS])
        for index, column in enumerate(_COLUMNS, start=1):
            values = [len(str(row[column])) for row in rows]
            sheet.column_dimensions[_excel_column(index)].width = min(
                max([len(column), *values]) + 2, 32
            )
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _pdf(rows: list[dict]) -> bytes:
        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=1 * cm,
            rightMargin=1 * cm,
            topMargin=1 * cm,
            bottomMargin=1 * cm,
            title="Telemetría ambiental IoT DHT22",
        )
        styles = getSampleStyleSheet()
        story = [Paragraph("Telemetría ambiental IoT DHT22", styles["Title"]), Spacer(1, 0.3 * cm)]
        headers = ["Fecha", "Dispositivo", "Temp °C", "Hum %", "Estado T", "Estado H", "Calidad"]
        data = [headers]
        for row in rows:
            data.append(
                [
                    row["measured_at"],
                    row["device_id"],
                    f'{row["temperature_celsius"]:.1f}',
                    f'{row["humidity_percent"]:.1f}',
                    row["temp_status"],
                    row["humidity_status"],
                    row["quality_flags"] or "OK",
                ]
            )
        if len(data) == 1:
            story.append(Paragraph("No hay muestras para exportar.", styles["BodyText"]))
        else:
            table = Table(data, repeatRows=1, colWidths=[4.4 * cm, 4 * cm, 2.3 * cm, 2.2 * cm, 2.3 * cm, 2.1 * cm, 6 * cm])
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 7),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#E8F5E9")]),
                    ]
                )
            )
            story.append(table)
        document.build(story)
        return buffer.getvalue()


def _to_row(sample: TelemetrySample) -> dict:
    return {
        "id": sample.id,
        "group_id": sample.group_id,
        "device_id": sample.device_id,
        "sequence": sample.sequence,
        "measured_at": _iso(sample.measured_at),
        "received_at": _iso(sample.received_at),
        "sampling_period_seconds": sample.sampling_period_seconds,
        "temperature_celsius": sample.ambient.temperature_celsius,
        "humidity_percent": sample.ambient.humidity_percent,
        "temp_status": sample.temp_status,
        "humidity_status": sample.humidity_status,
        "quality_flags": ",".join(sample.quality_flags),
    }


def _iso(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _excel_column(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result

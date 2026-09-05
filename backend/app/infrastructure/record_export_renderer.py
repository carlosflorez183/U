"""Renderizador de archivos CSV, JSON, XLSX y PDF para registros."""

from __future__ import annotations

import csv
import html
import io
import json
from datetime import datetime, timezone
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, XPreformatted

from app.application.exports import ExportedFile, ExportFormat
from app.domain.entities import Record


class FileRecordExportRenderer:
    """Adaptador que genera formatos descargables sin alterar entidades de dominio."""

    _MEDIA_TYPES = {
        ExportFormat.CSV: "text/csv; charset=utf-8",
        ExportFormat.JSON: "application/json; charset=utf-8",
        ExportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ExportFormat.PDF: "application/pdf",
    }
    _BASE_COLUMNS = ("id", "source", "created_at", "data_json")

    def render(self, records: Sequence[Record], export_format: ExportFormat) -> ExportedFile:
        filename = self._filename(export_format)
        if export_format is ExportFormat.JSON:
            content = self._render_json(records)
        elif export_format is ExportFormat.CSV:
            content = self._render_csv(records)
        elif export_format is ExportFormat.XLSX:
            content = self._render_xlsx(records)
        elif export_format is ExportFormat.PDF:
            content = self._render_pdf(records)
        else:  # Protección adicional si un adaptador llama el puerto sin FastAPI.
            raise ValueError(f"Formato de exportación no soportado: {export_format}")
        return ExportedFile(content=content, filename=filename, media_type=self._MEDIA_TYPES[export_format])

    def _render_json(self, records: Sequence[Record]) -> bytes:
        return json.dumps(
            [self._public_record(record) for record in records],
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        ).encode("utf-8")

    def _render_csv(self, records: Sequence[Record]) -> bytes:
        rows, columns = self._tabular_rows(records)
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
        # BOM para que Excel detecte UTF-8 al abrir el CSV directamente.
        return buffer.getvalue().encode("utf-8-sig")

    def _render_xlsx(self, records: Sequence[Record]) -> bytes:
        rows, columns = self._tabular_rows(records)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Registros"
        worksheet.append(columns)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{self._column_letter(len(columns))}{max(1, len(rows) + 1)}"

        for row in rows:
            worksheet.append([row.get(column, "") for column in columns])
        for index, column in enumerate(columns, start=1):
            longest = max((len(str(row.get(column, ""))) for row in rows), default=len(column))
            worksheet.column_dimensions[self._column_letter(index)].width = min(max(longest, len(column)) + 2, 45)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _render_pdf(self, records: Sequence[Record]) -> bytes:
        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=1.5 * cm,
            leftMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
            title="Exportación de registros",
        )
        styles = getSampleStyleSheet()
        heading = ParagraphStyle("RecordHeading", parent=styles["Heading2"], spaceAfter=8)
        metadata = ParagraphStyle("RecordMetadata", parent=styles["BodyText"], leading=14, spaceAfter=8)
        content = ParagraphStyle(
            "RecordContent",
            parent=styles["Code"],
            fontName="Courier",
            fontSize=7.5,
            leading=9,
            alignment=TA_LEFT,
        )
        story: list[Any] = [Paragraph("Exportación de registros", styles["Title"]), Spacer(1, 0.25 * cm)]
        if not records:
            story.append(Paragraph("No hay registros para exportar.", styles["BodyText"]))
        for number, record in enumerate(records, start=1):
            if number > 1:
                story.append(PageBreak())
            public_record = self._public_record(record)
            story.append(Paragraph(f"Registro {number}", heading))
            source = html.escape(str(public_record["source"] or "Sin origen"))
            created_at = html.escape(str(public_record["created_at"]))
            record_id = html.escape(str(public_record["id"]))
            story.append(Paragraph(f"<b>ID:</b> {record_id}<br/><b>Origen:</b> {source}<br/><b>Fecha:</b> {created_at}", metadata))
            formatted_json = json.dumps(public_record["data"], ensure_ascii=False, indent=2, allow_nan=False)
            story.append(XPreformatted(html.escape(formatted_json), content))
        document.build(story)
        return buffer.getvalue()

    def _tabular_rows(self, records: Sequence[Record]) -> tuple[list[dict[str, str]], list[str]]:
        rows: list[dict[str, str]] = []
        dynamic_columns: set[str] = set()
        for record in records:
            public_record = self._public_record(record)
            row = {
                "id": str(public_record["id"]),
                "source": self._safe_cell(public_record["source"]),
                "created_at": str(public_record["created_at"]),
                "data_json": self._safe_cell(json.dumps(public_record["data"], ensure_ascii=False, separators=(",", ":"), allow_nan=False)),
            }
            for column, value in self._flatten_data(public_record["data"]).items():
                dynamic_columns.add(column)
                row[column] = self._safe_cell(value)
            rows.append(row)
        return rows, [*self._BASE_COLUMNS, *sorted(dynamic_columns)]

    def _public_record(self, record: Record) -> dict[str, Any]:
        created_at = record.created_at.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
        return {"id": record.id, "data": dict(record.data), "source": record.source, "created_at": created_at}

    def _flatten_data(self, data: Any, prefix: str = "data") -> dict[str, str]:
        if isinstance(data, dict):
            flattened: dict[str, str] = {}
            for key in sorted(data, key=str):
                escaped_key = str(key).replace("~", "~0").replace("/", "~1")
                flattened.update(self._flatten_data(data[key], f"{prefix}/{escaped_key}"))
            return flattened
        if isinstance(data, list):
            return {prefix: json.dumps(data, ensure_ascii=False, separators=(",", ":"), allow_nan=False)}
        if data is None:
            return {prefix: ""}
        return {prefix: str(data)}

    @staticmethod
    def _safe_cell(value: Any) -> str:
        text = "" if value is None else str(value)
        # Evita que Excel evalúe valores procedentes de usuarios como fórmulas.
        return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text

    @staticmethod
    def _column_letter(index: int) -> str:
        letters = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            letters = chr(65 + remainder) + letters
        return letters

    @staticmethod
    def _filename(export_format: ExportFormat) -> str:
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        return f"registros_{timestamp}.{export_format.value}"
